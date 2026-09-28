import os
import sys
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

# Color Palette (Pink / Burgundy / Slate Banking Theme)
COLOR_PRIMARY = RGBColor(159, 28, 77)      # Deep Rose / Burgundy (#9F1C4D)
COLOR_SECONDARY = RGBColor(225, 29, 72)    # Rose Accent (#E11D48)
COLOR_DARK = RGBColor(15, 23, 42)          # Slate Charcoal (#0F172A)
COLOR_MUTED = RGBColor(100, 116, 139)      # Slate Gray (#64748B)
HEX_PRIMARY = "9F1C4D"
HEX_ACCENT = "E11D48"
HEX_LIGHT_BG = "FDF2F8"
HEX_BORDER = "FBCFE8"
HEX_HEADER_BG = "831843"

def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=120, bottom=120, left=150, right=150):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def add_styled_heading(doc, text, level):
    h = doc.add_heading(text, level=level)
    h.paragraph_format.space_before = Pt(14)
    h.paragraph_format.space_after = Pt(6)
    for run in h.runs:
        if level == 1:
            run.font.color.rgb = COLOR_PRIMARY
            run.font.size = Pt(18)
            run.font.bold = True
        elif level == 2:
            run.font.color.rgb = COLOR_SECONDARY
            run.font.size = Pt(14)
            run.font.bold = True
        elif level == 3:
            run.font.color.rgb = COLOR_DARK
            run.font.size = Pt(12)
            run.font.bold = True
    return h

def add_callout(doc, text, title="NOTE / ARCHITECTURAL PRINCIPLE"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_background(cell, HEX_LIGHT_BG)
    set_cell_margins(cell, top=140, bottom=140, left=200, right=200)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(4)
    r_title = p.add_run(f"📌 {title}\n")
    r_title.bold = True
    r_title.font.color.rgb = COLOR_PRIMARY
    r_title.font.size = Pt(10.5)
    
    r_body = p.add_run(text)
    r_body.font.color.rgb = COLOR_DARK
    r_body.font.size = Pt(10)
    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_after = Pt(6)

def format_table(tbl, col_widths, headers, rows_data):
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr_cells = tbl.rows[0].cells
    for i, header_text in enumerate(headers):
        hdr_cells[i].text = header_text
        set_cell_background(hdr_cells[i], HEX_HEADER_BG)
        set_cell_margins(hdr_cells[i], top=120, bottom=120, left=120, right=120)
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        for run in p.runs:
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.font.size = Pt(9.5)
            
    for row_idx, row_data in enumerate(rows_data):
        row = tbl.add_row()
        bg = "FFFFFF" if row_idx % 2 == 0 else "FDF2F8"
        for col_idx, text in enumerate(row_data):
            cell = row.cells[col_idx]
            cell.text = str(text)
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=90, bottom=90, left=110, right=110)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            for run in p.runs:
                run.font.size = Pt(9)
                run.font.color.rgb = COLOR_DARK

    for row in tbl.rows:
        for i, w in enumerate(col_widths):
            row.cells[i].width = Inches(w)

def create_title_page(doc, title, subtitle, authors, date_str):
    p_pre = doc.add_paragraph()
    p_pre.paragraph_format.space_before = Pt(40)
    
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_title = p_title.add_run(title)
    r_title.bold = True
    r_title.font.size = Pt(26)
    r_title.font.color.rgb = COLOR_PRIMARY
    
    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_after = Pt(30)
    r_sub = p_sub.add_run(subtitle)
    r_sub.font.size = Pt(14)
    r_sub.font.color.rgb = COLOR_SECONDARY
    
    # Divider
    p_div = doc.add_paragraph()
    p_div.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_div = p_div.add_run("―" * 35)
    r_div.font.color.rgb = COLOR_MUTED
    
    p_meta = doc.add_paragraph()
    p_meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_meta.paragraph_format.space_before = Pt(40)
    
    r_deg = p_meta.add_run("Full-Stack Software Engineering Capstone Defense Project\n")
    r_deg.bold = True
    r_deg.font.size = Pt(12)
    r_deg.font.color.rgb = COLOR_DARK
    
    r_auth = p_meta.add_run(f"Project Team: {', '.join(authors)}\n")
    r_auth.font.size = Pt(11)
    r_auth.font.color.rgb = COLOR_DARK
    
    r_date = p_meta.add_run(f"Date: {date_str}\nLocation: Metro Manila, Philippines\nCurrencies: PHP (₱) Standard")
    r_date.font.size = Pt(10.5)
    r_date.font.color.rgb = COLOR_MUTED
    
    doc.add_page_break()

print("Base Word doc builder initialized successfully.")
